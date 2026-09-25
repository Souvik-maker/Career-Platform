from pypdf import PdfReader
import io
import docx
import streamlit as st

def extract_text_from_file(uploaded_file) -> str:
    """Extracts raw text from uploaded PDF or DOCX files."""
    file_name = uploaded_file.name.lower()
    
    try:
        # 1. Handle PDF
        if file_name.endswith('.pdf'):
            pdf_reader = pypdf.PdfReader(uploaded_file)
            text = ""
            for page in pdf_reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
            return text.strip()

        # 2. Handle DOCX
        elif file_name.endswith('.docx'):
            doc = docx.Document(uploaded_file)
            full_text = []
            
            # Extract text from standard paragraphs
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text)
                    
            # Extract text from tables inside the docx (common in resumes)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            full_text.append(cell.text)

            return "\n".join(full_text).strip()

        # 3. Handle unsupported legacy format (.doc)
        elif file_name.endswith('.doc'):
            st.warning("Legacy .doc files are not fully supported. Please save/convert your file as .docx or .pdf.")
            return ""

        else:
            st.error("Unsupported file format. Please upload a PDF or DOCX file.")
            return ""

    except Exception as e:
        st.error(f"Error parsing resume file: {str(e)}")
        return ""